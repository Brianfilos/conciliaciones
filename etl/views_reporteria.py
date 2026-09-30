import re

from django.contrib.auth.decorators import login_required
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils.decorators import method_decorator
from django.views import View

from accounts import permisos
from muni.models import Municipio
from .models import Proceso
from .services import export_completo, reporteria, unicos


def _municipio(request):
    """Municipio a mostrar: el del usuario; un superusuario puede elegir con ?municipio=."""
    u = request.user
    if u.is_superuser or u.municipio_id is None:
        activos = Municipio.objects.filter(activo=True).order_by("orden", "nombre")
        return (activos.filter(codigo=request.GET.get("municipio", "")).first()
                or u.municipio or activos.first())
    return u.municipio


def _filtros(request):
    g = request.GET
    proceso = g.get("proceso", "")
    pago = g.get("pago", "").upper()
    periodo = g.get("periodo", "")
    return {
        "proceso": int(proceso) if proceso.isdigit() else None,
        "ano": g.get("ano") if re.fullmatch(r"\d{4}", g.get("ano", "")) else None,
        "pago": pago if pago in ("PAGADO", "PENDIENTE") else None,
        "cxc": g.get("cxc", "").strip().upper()[:40] or None,
        "periodo": periodo if re.fullmatch(r"\d{4}(-(\d{2}|B\d))?", periodo) else None,
        "doc": g.get("doc", "").strip()[:30] or None,
        "q": g.get("q", "").strip()[:60] or None,
    }


@method_decorator(login_required, name="dispatch")
class ReporteriaView(View):
    def get(self, request):
        municipio = _municipio(request)
        if municipio is None:
            return render(request, "etl/reporteria.html", {"datos": None, "municipios": []})
        datos = reporteria.calcular(municipio, _filtros(request))
        municipios = []
        if request.user.is_superuser or request.user.municipio_id is None:
            municipios = list(Municipio.objects.filter(activo=True).order_by("orden", "nombre")
                              .values("codigo", "nombre"))
        return render(request, "etl/reporteria.html", {
            "datos": datos, "municipios": municipios, "municipio_actual": municipio,
        })


@method_decorator(login_required, name="dispatch")
class ReporteriaDatosView(View):
    def get(self, request):
        municipio = _municipio(request)
        if municipio is None:
            return JsonResponse({"error": "Sin municipio"}, status=404)
        return JsonResponse(reporteria.calcular(municipio, _filtros(request)))


@method_decorator(login_required, name="dispatch")
class ReporteUnicosView(View):
    """Descarga las declaraciones "únicas" del proceso (una por vigencia, ver services/unicos.py)."""

    def get(self, request, proceso_id):
        proceso = get_object_or_404(Proceso, id=proceso_id)
        if not permisos.del_municipio(request.user, proceso.municipio):
            return redirect("reporteria")

        import openpyxl
        from openpyxl.styles import Font, PatternFill

        headers, filas_unicas, filas_excluidas = unicos.build_rows(proceso)

        wb = openpyxl.Workbook()
        ws1 = wb.active
        ws1.title = "unicos"
        ws1.append(headers)
        ws2 = wb.create_sheet("excluidas_duplicadas")
        ws2.append(headers)
        for ws, filas in ((ws1, filas_unicas), (ws2, filas_excluidas)):
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1B3A6B")
            for fila in filas:
                ws.append(["" if v is None else v for v in fila])

        import io
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        resp = HttpResponse(
            buf.read(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        nombre = f"{proceso.municipio.codigo}_{proceso.codigo}_unicos.xlsx"
        resp["Content-Disposition"] = f'attachment; filename="{nombre}"'
        return resp


@method_decorator(login_required, name="dispatch")
class ExportarCompletoView(View):
    """Export con las columnas crudas de GOBS (ver services/export_completo.py), para las
    declaraciones que cumplan pago/año — con selector de columnas y de si incluir
    duplicados. Primero muestra el selector; con ?descargar=1 genera el archivo."""

    def _procesos_con_gobs(self, municipio):
        from etl.services import gobs_pg
        return [p for p in Proceso.objects.filter(municipio=municipio, activo=True).order_by("orden")
                if gobs_pg.fuente_para(municipio.codigo, p.codigo) is not None]

    def get(self, request):
        municipio = _municipio(request)
        if municipio is None:
            return redirect("reporteria")
        procesos = self._procesos_con_gobs(municipio)
        if not procesos:
            return redirect("reporteria")

        proceso_id = request.GET.get("proceso", "")
        proceso = next((p for p in procesos if str(p.id) == proceso_id), procesos[0])
        if not permisos.del_municipio(request.user, proceso.municipio):
            return redirect("reporteria")

        pago = request.GET.get("pago", "").upper()
        pago = pago if pago in ("PAGADO", "PENDIENTE") else ""
        ano = request.GET.get("ano", "") if re.fullmatch(r"\d{4}", request.GET.get("ano", "")) else ""
        incluir_duplicados = request.GET.get("dup") == "1"
        filtros = {"pago": pago or None, "ano": ano or None}

        disponibles = export_completo.columnas_disponibles(municipio.codigo, proceso.codigo)

        if request.GET.get("descargar") == "1":
            elegidas = [c for c in request.GET.getlist("col") if c in disponibles]
            headers, filas = export_completo.generar(proceso, filtros, elegidas, incluir_duplicados)
            if not headers:
                messages_txt = "No hay datos de GOBS para exportar con estos filtros."
                return render(request, "etl/exportar_completo.html", {
                    "municipio": municipio, "procesos": procesos, "proceso": proceso,
                    "disponibles": disponibles, "elegidas": elegidas or disponibles,
                    "pago": pago, "ano": ano, "incluir_duplicados": incluir_duplicados,
                    "error": messages_txt,
                })
            import io
            import openpyxl
            from openpyxl.styles import Font, PatternFill
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "datos"
            ws.append(headers)
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1B3A6B")
            for fila in filas:
                ws.append(["" if v is None or (isinstance(v, float) and v != v) else v for v in fila])
            buf = io.BytesIO()
            wb.save(buf)
            buf.seek(0)
            resp = HttpResponse(
                buf.read(),
                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            sufijo = "_con_duplicados" if incluir_duplicados else ""
            nombre = f"{municipio.codigo}_{proceso.codigo}_completo{sufijo}.xlsx"
            resp["Content-Disposition"] = f'attachment; filename="{nombre}"'
            return resp

        return render(request, "etl/exportar_completo.html", {
            "municipio": municipio, "procesos": procesos, "proceso": proceso,
            "disponibles": disponibles, "elegidas": disponibles,
            "pago": pago, "ano": ano, "incluir_duplicados": incluir_duplicados,
        })


@method_decorator(login_required, name="dispatch")
class ReporteriaBuscarView(View):
    """Sugerencias del buscador de contribuyentes."""
    def get(self, request):
        municipio = _municipio(request)
        if municipio is None:
            return JsonResponse({"resultados": []})
        return JsonResponse({"resultados": reporteria.buscar(municipio, request.GET.get("q", ""))})
