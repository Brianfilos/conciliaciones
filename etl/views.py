import io
import csv
import threading
from urllib.parse import urlencode
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.utils.decorators import method_decorator
from django.views import View
from django.core.paginator import Paginator
from django.http import HttpResponse, JsonResponse
from django.contrib import messages
from django.utils import timezone
from .models import Proceso, Ejecucion, InsumoEjecucion, EncabezadoCXC, DetalleCXC
from .forms import EjecutarProcesoForm
from .services.motor import MotorETL


def _parse_fecha(valor):
    """Acepta YYYY-MM-DD (input type=date) o DD-MM-YYYY escrito manualmente.
    Retorna datetime.date o None."""
    from datetime import date as _date, datetime as _dt
    v = (valor or "").strip()
    if not v:
        return None
    try:
        if len(v) == 10 and v[2] == "-" and v[5] == "-":
            # DD-MM-YYYY
            d, m, a = v.split("-")
            return _date(int(a), int(m), int(d))
        # YYYY-MM-DD
        return _dt.strptime(v, "%Y-%m-%d").date()
    except Exception:
        return None


def _run_motor(ejecucion_id, archivos, filtros=None):
    """Ejecuta el motor ETL en un hilo separado."""
    import django
    from etl.models import Ejecucion as Ej
    from etl.services.motor import MotorETL as M
    try:
        ejecucion = Ej.objects.get(id=ejecucion_id)
        motor = M(ejecucion)
        motor.ejecutar(archivos, filtros)
    except Exception:
        pass


@method_decorator(login_required, name="dispatch")
class ProcesoListView(View):
    def get(self, request):
        return redirect("municipio_home")


@method_decorator(login_required, name="dispatch")
class EjecutarProcesoView(View):
    def _get_proceso(self, proceso_id, request):
        proceso = get_object_or_404(Proceso, id=proceso_id)
        if proceso.municipio != request.user.municipio and not request.user.is_admin:
            return None
        return proceso

    def get(self, request, proceso_id):
        proceso = self._get_proceso(proceso_id, request)
        if not proceso:
            return redirect("municipio_home")
        form = EjecutarProcesoForm(proceso)
        ultima = Ejecucion.objects.filter(proceso=proceso).first()
        return render(request, "etl/ejecutar.html", {"proceso": proceso, "form": form, "ultima": ultima})

    def post(self, request, proceso_id):
        proceso = self._get_proceso(proceso_id, request)
        if not proceso:
            return redirect("municipio_home")
        form = EjecutarProcesoForm(proceso, request.POST, request.FILES)
        if not form.is_valid():
            ultima = Ejecucion.objects.filter(proceso=proceso).first()
            return render(request, "etl/ejecutar.html", {"proceso": proceso, "form": form, "ultima": ultima})

        ejecucion = Ejecucion.objects.create(proceso=proceso, usuario=request.user)
        archivos = {}
        for insumo in proceso.insumos.filter(tipo="CARGUE"):
            if insumo.nombre_campo in request.FILES:
                archivo = request.FILES[insumo.nombre_campo]
                ins_ej = InsumoEjecucion.objects.create(
                    ejecucion=ejecucion,
                    insumo_def=insumo,
                    archivo=archivo,
                    nombre_original=archivo.name,
                )
                archivos[insumo.nombre_campo] = ins_ej.archivo.path

        # Lanzar en background y redirigir a pantalla de progreso
        filtros = {"desde": form.cleaned_data.get("fecha_desde"),
                   "hasta": form.cleaned_data.get("fecha_hasta")}
        t = threading.Thread(target=_run_motor, args=(ejecucion.id, archivos, filtros), daemon=True)
        t.start()
        return redirect("ejecutar_progreso", ejecucion_id=ejecucion.id)


@method_decorator(login_required, name="dispatch")
class LimpiarProcesoView(View):
    """Borra todos los encabezados, detalles y ejecuciones de un proceso. Solo admins."""

    def post(self, request, proceso_id):
        if not request.user.is_admin:
            messages.error(request, "No tienes permisos para realizar esta acción.")
            return redirect("ejecutar_proceso", proceso_id=proceso_id)

        proceso = get_object_or_404(Proceso, id=proceso_id)
        # Verificar que el proceso pertenece al municipio del usuario (o es superusuario)
        if proceso.municipio != request.user.municipio and not request.user.is_superuser:
            messages.error(request, "No tienes permisos sobre este proceso.")
            return redirect("ejecutar_proceso", proceso_id=proceso_id)

        enc_count = proceso.encabezados.count()
        eje_count = proceso.ejecuciones.count()
        proceso.encabezados.all().delete()
        proceso.ejecuciones.all().delete()

        messages.success(
            request,
            f"Datos limpiados: {enc_count} encabezados y {eje_count} ejecuciones eliminados del proceso «{proceso.nombre}»."
        )
        return redirect("admin_municipio")


@method_decorator(login_required, name="dispatch")
class HistorialView(View):
    def get(self, request):
        municipio = request.user.municipio
        qs = Ejecucion.objects.filter(proceso__municipio=municipio).select_related("proceso", "usuario")
        paginator = Paginator(qs, 20)
        page_obj = paginator.get_page(request.GET.get("page", 1))
        return render(request, "etl/historial.html", {"page_obj": page_obj, "municipio": municipio})


@login_required
def ejecutar_progreso(request, ejecucion_id):
    ejecucion = get_object_or_404(Ejecucion, id=ejecucion_id)
    if ejecucion.proceso.municipio != request.user.municipio and not request.user.is_admin:
        return redirect("municipio_home")
    return render(request, "etl/progreso.html", {"ejecucion": ejecucion})


@login_required
def ejecucion_status(request, ejecucion_id):
    ejecucion = get_object_or_404(Ejecucion, id=ejecucion_id)
    return JsonResponse({
        "estado": ejecucion.estado,
        "nuevos": ejecucion.registros_nuevos,
        "duplicados": ejecucion.registros_duplicados,
        "error": ejecucion.error_log[:400] if ejecucion.error_log else "",
        "proceso_id": ejecucion.proceso_id,
    })


@method_decorator(login_required, name="dispatch")
class DashboardView(View):
    def get(self, request, proceso_id):
        proceso = get_object_or_404(Proceso, id=proceso_id)
        if proceso.municipio != request.user.municipio and not request.user.is_admin:
            return redirect("municipio_home")

        es_caldas     = proceso.municipio.codigo == "CALDAS"
        es_envigado   = proceso.municipio.codigo == "ENVIGADO"
        es_caldas_ica = es_caldas and proceso.codigo == "DECLAREYPAGUE"

        tab = request.GET.get("tab", "encabezado")
        q = request.GET.get("q", "").strip()
        # column text filters
        q_consec        = request.GET.get("q_consec", "").strip()
        q_num           = request.GET.get("q_num", "").strip()
        q_nombre        = request.GET.get("q_nombre", "").strip()
        q_desc          = request.GET.get("q_desc", "").strip()
        q_concepto      = request.GET.get("q_concepto", "").strip()
        q_centro        = request.GET.get("q_centro", "").strip()
        # column value-list filters
        tipo_doc_filtro    = request.GET.get("tipo_doc", "").strip()
        estado_pago_filtro = request.GET.get("estado_pago", "").strip()
        estado_filtro      = request.GET.get("estado", "").strip()
        # Caldas extra filters
        caldas_tipo_persona_filtro = request.GET.get("caldas_tipo_persona", "").strip()
        caldas_periodo_filtro      = request.GET.get("caldas_periodo", "").strip()
        caldas_ano_filtro          = request.GET.get("caldas_ano", "").strip()
        caldas_clasi_filtro        = request.GET.get("caldas_clasi", "").strip()
        # Envigado extra filters (rango desde/hasta)
        env_fecha_pago_desde = _parse_fecha(request.GET.get("env_fecha_pago_desde", ""))
        env_fecha_pago_hasta = _parse_fecha(request.GET.get("env_fecha_pago_hasta", ""))
        env_fecha_pres_desde = _parse_fecha(request.GET.get("env_fecha_pres_desde", ""))
        env_fecha_pres_hasta = _parse_fecha(request.GET.get("env_fecha_pres_hasta", ""))
        q_valor_unit  = request.GET.get("q_valor_unit", "").strip()
        q_valor_tot   = request.GET.get("q_valor_tot", "").strip()
        q_total_pagar = request.GET.get("q_total_pagar", "").strip()
        fecha_desde = _parse_fecha(request.GET.get("fecha_desde", ""))
        fecha_hasta = _parse_fecha(request.GET.get("fecha_hasta", ""))
        per_page = int(request.GET.get("per_page", 20))
        if per_page not in [20, 50, 100]:
            per_page = 20

        if tab == "detalle":
            qs = DetalleCXC.objects.filter(encabezado__proceso=proceso).select_related("encabezado")
            # Aplicar todos los filtros de encabezado al detalle (filtro inteligente cruzado)
            if q:
                qs = (qs.filter(encabezado__consecutivo_cxc__icontains=q) |
                      qs.filter(encabezado__numero_documento__icontains=q) |
                      qs.filter(encabezado__razon_social__icontains=q) |
                      qs.filter(encabezado__primer_apellido__icontains=q) |
                      qs.filter(encabezado__segundo_apellido__icontains=q) |
                      qs.filter(encabezado__primer_nombre__icontains=q) |
                      qs.filter(encabezado__descripcion__icontains=q) |
                      qs.filter(codigo_concepto__icontains=q) |
                      qs.filter(centro_costo__icontains=q))
            if q_consec:
                qs = qs.filter(encabezado__consecutivo_cxc__icontains=q_consec)
            if q_num:
                qs = qs.filter(encabezado__numero_documento__icontains=q_num)
            if q_nombre:
                qs = qs.filter(encabezado__primer_apellido__icontains=q_nombre) | \
                     qs.filter(encabezado__razon_social__icontains=q_nombre)
            if q_concepto:
                qs = qs.filter(codigo_concepto__icontains=q_concepto)
            if q_centro:
                qs = qs.filter(centro_costo__icontains=q_centro)
            if tipo_doc_filtro:
                qs = qs.filter(encabezado__tipo_documento__iexact=tipo_doc_filtro)
            if estado_pago_filtro:
                qs = qs.filter(encabezado__estado_pago__icontains=estado_pago_filtro)
            if fecha_desde:
                qs = qs.filter(encabezado__fecha_cobro__gte=fecha_desde)
            if fecha_hasta:
                qs = qs.filter(encabezado__fecha_cobro__lte=fecha_hasta)
            if estado_filtro == "__blank__":
                qs = qs.filter(encabezado__estado_cxc="")
            elif estado_filtro:
                qs = qs.filter(encabezado__estado_cxc__iexact=estado_filtro)
            if caldas_tipo_persona_filtro:
                qs = qs.filter(encabezado__datos_extra__tipo_persona=caldas_tipo_persona_filtro)
            if caldas_periodo_filtro:
                qs = qs.filter(encabezado__datos_extra__periodo=caldas_periodo_filtro)
            if caldas_ano_filtro:
                qs = qs.filter(encabezado__datos_extra__ano=caldas_ano_filtro)
            if caldas_clasi_filtro:
                qs = qs.filter(centro_costo=caldas_clasi_filtro)
            if q_valor_unit:
                from django.db.models.functions import Cast as _Cast
                from django.db.models import CharField as _CF
                qs = qs.annotate(_vu_s=_Cast("valor_unitario", _CF())).filter(_vu_s__icontains=q_valor_unit)
            if q_valor_tot:
                from django.db.models.functions import Cast as _Cast
                from django.db.models import CharField as _CF
                qs = qs.annotate(_vt_s=_Cast("valor_total", _CF())).filter(_vt_s__icontains=q_valor_tot)
            # Envigado: filtros de fecha_pago / fecha_pres cruzados al detalle
            from datetime import datetime as _dt2
            for _campo, _desde, _hasta in [
                ("fecha_pago", env_fecha_pago_desde, env_fecha_pago_hasta),
                ("fecha_pres", env_fecha_pres_desde, env_fecha_pres_hasta),
            ]:
                if _desde or _hasta:
                    _vals = []
                    for _v in EncabezadoCXC.objects.filter(proceso=proceso)\
                            .values_list(f"datos_extra__{_campo}", flat=True).distinct():
                        if not _v:
                            continue
                        try:
                            _d = _dt2.strptime(str(_v), "%d-%m-%Y").date()
                        except ValueError:
                            continue
                        if _desde and _d < _desde:
                            continue
                        if _hasta and _d > _hasta:
                            continue
                        _vals.append(_v)
                    qs = qs.filter(**{f"encabezado__datos_extra__{_campo}__in": _vals})
            total_sin_filtro = DetalleCXC.objects.filter(encabezado__proceso=proceso).count()
            # unique values for detalle dropdowns
            codigos_concepto = list(DetalleCXC.objects.filter(encabezado__proceso=proceso)
                                    .values_list("codigo_concepto", flat=True).distinct().order_by("codigo_concepto"))
            centros_costo = list(DetalleCXC.objects.filter(encabezado__proceso=proceso)
                                 .values_list("centro_costo", flat=True).distinct().order_by("centro_costo"))
            tipos_documento = []; estados_pago = []; estados = []
        else:
            tab = "encabezado"
            qs = EncabezadoCXC.objects.filter(proceso=proceso)
            if q:
                qs = qs.filter(consecutivo_cxc__icontains=q) | \
                     qs.filter(numero_documento__icontains=q) | \
                     qs.filter(primer_apellido__icontains=q) | \
                     qs.filter(razon_social__icontains=q) | \
                     qs.filter(descripcion__icontains=q)
            if q_consec:
                qs = qs.filter(consecutivo_cxc__icontains=q_consec)
            if q_num:
                qs = qs.filter(numero_documento__icontains=q_num)
            if q_nombre:
                qs = qs.filter(primer_apellido__icontains=q_nombre) | qs.filter(razon_social__icontains=q_nombre)
            if q_desc:
                qs = qs.filter(descripcion__icontains=q_desc)
            if tipo_doc_filtro:
                qs = qs.filter(tipo_documento__iexact=tipo_doc_filtro)
            if estado_pago_filtro:
                qs = qs.filter(estado_pago__icontains=estado_pago_filtro)
            if fecha_desde:
                qs = qs.filter(fecha_cobro__gte=fecha_desde)
            if fecha_hasta:
                qs = qs.filter(fecha_cobro__lte=fecha_hasta)
            if estado_filtro == "__blank__":
                qs = qs.filter(estado_cxc="")
            elif estado_filtro:
                qs = qs.filter(estado_cxc__iexact=estado_filtro)
            # Caldas: filtros adicionales por datos_extra (tipo_persona, periodo, año)
            if caldas_tipo_persona_filtro:
                qs = qs.filter(datos_extra__tipo_persona=caldas_tipo_persona_filtro)
            if caldas_periodo_filtro:
                qs = qs.filter(datos_extra__periodo=caldas_periodo_filtro)
            if caldas_ano_filtro:
                qs = qs.filter(datos_extra__ano=caldas_ano_filtro)
            # Envigado: filtros por rango en fecha_pago y fecha_pres (strings "dd-mm-yyyy" en datos_extra)
            for campo, desde, hasta in [
                ("fecha_pago", env_fecha_pago_desde, env_fecha_pago_hasta),
                ("fecha_pres", env_fecha_pres_desde, env_fecha_pres_hasta),
            ]:
                if desde or hasta:
                    from datetime import datetime
                    vals_en_rango = []
                    for v in EncabezadoCXC.objects.filter(proceso=proceso)\
                            .values_list(f"datos_extra__{campo}", flat=True).distinct():
                        if not v:
                            continue
                        try:
                            d = datetime.strptime(str(v), "%d-%m-%Y").date()
                        except ValueError:
                            continue
                        if desde and d < desde:
                            continue
                        if hasta and d > hasta:
                            continue
                        vals_en_rango.append(v)
                    qs = qs.filter(**{f"datos_extra__{campo}__in": vals_en_rango})
            if q_total_pagar:
                from django.db.models.functions import Cast as _Cast
                from django.db.models import CharField as _CF
                qs = qs.annotate(_tp_s=_Cast("total_a_pagar", _CF())).filter(_tp_s__icontains=q_total_pagar)

            total_sin_filtro = EncabezadoCXC.objects.filter(proceso=proceso).count()
            base = EncabezadoCXC.objects.filter(proceso=proceso)
            tipos_documento = list(base.values_list("tipo_documento", flat=True).distinct().order_by("tipo_documento"))
            estados_pago    = list(base.values_list("estado_pago", flat=True).distinct().order_by("estado_pago"))
            estados         = list(base.values_list("estado_cxc", flat=True).distinct().order_by("estado_cxc"))
            codigos_concepto = []; centros_costo = []

        # Caldas: valores únicos para filtros de lista
        caldas_periodos   = []
        caldas_anos       = []
        caldas_clasif     = []
        if (es_caldas or es_envigado) and tab == "encabezado":
            from django.db.models.functions import Cast
            from django.db.models import CharField
            caldas_periodos = sorted(set(
                v for v in EncabezadoCXC.objects.filter(proceso=proceso)
                    .values_list("datos_extra__periodo", flat=True) if v
            ))
            caldas_anos = sorted(set(
                v for v in EncabezadoCXC.objects.filter(proceso=proceso)
                    .values_list("datos_extra__ano", flat=True) if v
            ))
        if es_caldas and tab == "detalle":
            caldas_clasif = sorted(set(
                v for v in DetalleCXC.objects.filter(encabezado__proceso=proceso)
                    .values_list("centro_costo", flat=True) if v
            ))

        filtros_activos = len([x for x in [q, q_consec, q_num, q_nombre, q_desc,
                                            q_concepto, q_centro,
                                            q_valor_unit, q_valor_tot, q_total_pagar,
                                            tipo_doc_filtro, estado_pago_filtro,
                                            fecha_desde, fecha_hasta, estado_filtro,
                                            caldas_tipo_persona_filtro, caldas_periodo_filtro,
                                            caldas_ano_filtro, caldas_clasi_filtro,
                                            env_fecha_pago_desde, env_fecha_pago_hasta,
                                            env_fecha_pres_desde, env_fecha_pres_hasta] if x])
        total_filtrado = qs.count()
        paginator = Paginator(qs, per_page)
        try:
            page_num = int(request.GET.get("page") or 1)
        except (ValueError, TypeError):
            page_num = 1
        page_obj = paginator.get_page(max(1, min(page_num, paginator.num_pages)))

        # Para Caldas Pagos: resolver descripción SAIMYR de cada concepto
        caldas_ica_totals = {}
        caldas_desc_map = {}
        if es_caldas and tab == "detalle":
            from etl.services.exportador_caldas import _desc as caldas_desc_fn
            for det in page_obj.object_list:
                caldas_desc_map[det.pk] = caldas_desc_fn(
                    det.codigo_concepto, (det.centro_costo or "").strip()
                )

        # Para Caldas ICA Declaraciones: totales financieros por encabezado
        caldas_ica_totals = {}
        if es_caldas_ica and tab == "encabezado":
            OCC_ICA   = {"OCC-0048", "OCC-0047", "OCC-046"}
            OCC_RET   = {"OCC-993"}
            OCC_AUT   = {"OCC-209"}
            OCC_SAN   = {"OCC-04","OCC-05","OCC-06","OCC-07","OCC-08","OCC-09","OCC-10"}
            OCC_INT   = {"OCC-23"}
            OCC_BOM   = {"OCC-051"}
            OCC_AVI   = {"69"}
            from etl.models import DetalleCXC as _Det
            det_qs = _Det.objects.filter(
                encabezado__in=[obj.pk for obj in page_obj.object_list]
            ).values("encabezado_id", "codigo_concepto", "valor_total", "centro_costo")
            for d in det_qs:
                eid = d["encabezado_id"]
                c   = d["codigo_concepto"] or ""
                v   = float(d["valor_total"] or 0)
                t   = caldas_ica_totals.setdefault(eid, {
                    "iyc": 0, "retencion": 0, "autorretencion": 0,
                    "anticipo": 0, "sanciones": 0, "intereses": 0,
                    "bomberil": 0, "avisos": 0, "clase": set()
                })
                if c in OCC_ICA:
                    t["iyc"] += v
                    clasi = (d["centro_costo"] or "").strip()
                    if clasi:
                        t["clase"].add(clasi)
                elif c in OCC_RET:
                    t["retencion"] += abs(v)
                elif c in OCC_AUT:
                    t["autorretencion"] += abs(v)
                elif c in OCC_SAN:
                    t["sanciones"] += v
                elif c in OCC_INT:
                    t["intereses"] += v
                elif c in OCC_BOM:
                    t["bomberil"] += v
                elif c in OCC_AVI:
                    t["avisos"] += v
            # Convertir set de clases a string
            for t in caldas_ica_totals.values():
                t["clase"] = "/".join(sorted(t["clase"]))

        # Sufijo A/R para columna Retencion_autoretencion en export declaraciones Caldas
        caldas_ret_auto = ""
        if es_caldas:
            if "RETE" in proceso.codigo.upper():
                caldas_ret_auto = "R"
            elif not es_caldas_ica:
                caldas_ret_auto = "A"

        return render(request, "dashboard/explorer.html", {
            "proceso": proceso,
            "tab": tab,
            "es_envigado": es_envigado,
            "es_caldas": es_caldas,
            "es_caldas_ica": es_caldas_ica,
            "caldas_ret_auto": caldas_ret_auto,
            "tab_enc_label": "Declaraciones" if es_caldas else "Encabezado",
            "tab_det_label": "Pagos" if es_caldas else "Detalle",
            "page_obj": page_obj,
            "total_sin_filtro": total_sin_filtro,
            "total_filtrado": total_filtrado,
            "filtros_activos": filtros_activos,
            "q": q,
            "q_consec": q_consec, "q_num": q_num, "q_nombre": q_nombre,
            "q_desc": q_desc, "q_concepto": q_concepto, "q_centro": q_centro,
            "q_valor_unit": q_valor_unit, "q_valor_tot": q_valor_tot, "q_total_pagar": q_total_pagar,
            "export_filter_qs": urlencode([
                (k, v) for k, v in [
                    ("q", q), ("q_consec", q_consec), ("q_num", q_num), ("q_nombre", q_nombre),
                    ("q_desc", q_desc), ("q_concepto", q_concepto), ("q_centro", q_centro),
                    ("q_valor_unit", q_valor_unit), ("q_valor_tot", q_valor_tot), ("q_total_pagar", q_total_pagar),
                    ("tipo_doc", tipo_doc_filtro), ("estado_pago", estado_pago_filtro), ("estado", estado_filtro),
                    ("fecha_desde", fecha_desde.isoformat() if fecha_desde else ""),
                    ("fecha_hasta", fecha_hasta.isoformat() if fecha_hasta else ""),
                    ("caldas_tipo_persona", caldas_tipo_persona_filtro), ("caldas_periodo", caldas_periodo_filtro),
                    ("caldas_ano", caldas_ano_filtro), ("caldas_clasi", caldas_clasi_filtro),
                    ("env_fecha_pago_desde", env_fecha_pago_desde.isoformat() if env_fecha_pago_desde else ""),
                    ("env_fecha_pago_hasta", env_fecha_pago_hasta.isoformat() if env_fecha_pago_hasta else ""),
                    ("env_fecha_pres_desde", env_fecha_pres_desde.isoformat() if env_fecha_pres_desde else ""),
                    ("env_fecha_pres_hasta", env_fecha_pres_hasta.isoformat() if env_fecha_pres_hasta else ""),
                ] if v
            ]),
            "pagination_qs": urlencode([
                (k, v) for k, v in [
                    ("tab", tab), ("per_page", str(per_page)),
                    ("q", q), ("q_consec", q_consec), ("q_num", q_num), ("q_nombre", q_nombre),
                    ("q_desc", q_desc), ("q_concepto", q_concepto), ("q_centro", q_centro),
                    ("q_valor_unit", q_valor_unit), ("q_valor_tot", q_valor_tot), ("q_total_pagar", q_total_pagar),
                    ("tipo_doc", tipo_doc_filtro), ("estado_pago", estado_pago_filtro), ("estado", estado_filtro),
                    ("fecha_desde", fecha_desde.isoformat() if fecha_desde else ""),
                    ("fecha_hasta", fecha_hasta.isoformat() if fecha_hasta else ""),
                    ("caldas_tipo_persona", caldas_tipo_persona_filtro), ("caldas_periodo", caldas_periodo_filtro),
                    ("caldas_ano", caldas_ano_filtro), ("caldas_clasi", caldas_clasi_filtro),
                    ("env_fecha_pago_desde", env_fecha_pago_desde.isoformat() if env_fecha_pago_desde else ""),
                    ("env_fecha_pago_hasta", env_fecha_pago_hasta.isoformat() if env_fecha_pago_hasta else ""),
                    ("env_fecha_pres_desde", env_fecha_pres_desde.isoformat() if env_fecha_pres_desde else ""),
                    ("env_fecha_pres_hasta", env_fecha_pres_hasta.isoformat() if env_fecha_pres_hasta else ""),
                ] if v
            ]),
            "tipo_doc_filtro": tipo_doc_filtro,
            "estado_pago_filtro": estado_pago_filtro,
            "fecha_desde": fecha_desde.isoformat() if fecha_desde else "",
            "fecha_hasta": fecha_hasta.isoformat() if fecha_hasta else "",
            "estado_filtro": estado_filtro,
            "per_page": per_page,
            "tipos_documento": tipos_documento,
            "estados_pago": estados_pago,
            "estados": estados,
            "codigos_concepto": codigos_concepto,
            "centros_costo": centros_costo,
            "caldas_desc_map": caldas_desc_map,
            "caldas_ica_totals": caldas_ica_totals,
            "caldas_periodos": caldas_periodos,
            "caldas_anos": caldas_anos,
            "caldas_clasif": caldas_clasif,
            "caldas_tipo_persona_filtro": caldas_tipo_persona_filtro,
            "caldas_periodo_filtro": caldas_periodo_filtro,
            "caldas_ano_filtro": caldas_ano_filtro,
            "caldas_clasi_filtro": caldas_clasi_filtro,
            "env_fecha_pago_desde": env_fecha_pago_desde.isoformat() if env_fecha_pago_desde else "",
            "env_fecha_pago_hasta": env_fecha_pago_hasta.isoformat() if env_fecha_pago_hasta else "",
            "env_fecha_pres_desde": env_fecha_pres_desde.isoformat() if env_fecha_pres_desde else "",
            "env_fecha_pres_hasta": env_fecha_pres_hasta.isoformat() if env_fecha_pres_hasta else "",
        })


@method_decorator(login_required, name="dispatch")
class ExportarView(View):
    def get(self, request, proceso_id):
        proceso = get_object_or_404(Proceso, id=proceso_id)
        if proceso.municipio != request.user.municipio and not request.user.is_admin:
            return redirect("municipio_home")

        fmt = request.GET.get("format", "excel")
        tab = request.GET.get("tab", "encabezado")
        q             = request.GET.get("q", "").strip()
        q_consec      = request.GET.get("q_consec", "").strip()
        q_num         = request.GET.get("q_num", "").strip()
        q_nombre      = request.GET.get("q_nombre", "").strip()
        q_desc        = request.GET.get("q_desc", "").strip()
        q_concepto    = request.GET.get("q_concepto", "").strip()
        q_centro      = request.GET.get("q_centro", "").strip()
        q_valor_unit  = request.GET.get("q_valor_unit", "").strip()
        q_valor_tot   = request.GET.get("q_valor_tot", "").strip()
        q_total_pagar = request.GET.get("q_total_pagar", "").strip()
        tipo_doc_filtro    = request.GET.get("tipo_doc", "").strip()
        estado_pago_filtro = request.GET.get("estado_pago", "").strip()
        estado_filtro      = request.GET.get("estado", "").strip()
        fecha_desde = _parse_fecha(request.GET.get("fecha_desde", ""))
        fecha_hasta = _parse_fecha(request.GET.get("fecha_hasta", ""))
        caldas_tipo_persona_filtro = request.GET.get("caldas_tipo_persona", "").strip()
        caldas_periodo_filtro      = request.GET.get("caldas_periodo", "").strip()
        caldas_ano_filtro          = request.GET.get("caldas_ano", "").strip()
        caldas_clasi_filtro        = request.GET.get("caldas_clasi", "").strip()
        env_fecha_pago_desde = _parse_fecha(request.GET.get("env_fecha_pago_desde", ""))
        env_fecha_pago_hasta = _parse_fecha(request.GET.get("env_fecha_pago_hasta", ""))
        env_fecha_pres_desde = _parse_fecha(request.GET.get("env_fecha_pres_desde", ""))
        env_fecha_pres_hasta = _parse_fecha(request.GET.get("env_fecha_pres_hasta", ""))

        nombre_base = f"{proceso.municipio.codigo}_{proceso.codigo}"

        # ── Caldas: formato SAIMYR ────────────────────────────────────────────
        if proceso.municipio.codigo == "CALDAS":
            from etl.services.exportador_caldas import build_rows, build_dec_rows, HEADERS
            filtros = {
                "q": q, "fecha_desde": fecha_desde, "fecha_hasta": fecha_hasta,
                "estado": estado_filtro, "estado_pago": estado_pago_filtro,
                "q_consec": q_consec, "q_num": q_num, "q_nombre": q_nombre,
                "caldas_tipo_persona": caldas_tipo_persona_filtro,
                "caldas_periodo": caldas_periodo_filtro,
                "caldas_ano": caldas_ano_filtro,
                "caldas_clasi": caldas_clasi_filtro,
            }
            if tab == "encabezado":
                headers, data = build_dec_rows(proceso, filtros)
            else:
                headers, data = build_rows(proceso, filtros)

            if fmt == "excel":
                import openpyxl
                from openpyxl.styles import Font, PatternFill, Alignment
                wb = openpyxl.Workbook()
                ws = wb.active
                ws.title = "PAGOS"
                ws.append(headers)
                for cell in ws[1]:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="1a6b3c")
                    cell.alignment = Alignment(horizontal="center")
                for row in data:
                    ws.append(["" if v is None else v for v in row])
                # Formato numérico 0.00 (sin símbolo moneda) en columna valor_pagado
                if "valor_pagado" in headers:
                    vcol = headers.index("valor_pagado") + 1
                    for row_cells in ws.iter_rows(min_row=2, min_col=vcol, max_col=vcol):
                        for cell in row_cells:
                            cell.number_format = "0.00"
                buf = io.BytesIO()
                wb.save(buf)
                buf.seek(0)
                resp = HttpResponse(
                    buf.read(),
                    content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )
                resp["Content-Disposition"] = f'attachment; filename="{nombre_base}_saimyr.xlsx"'
                return resp

            if fmt == "csv":
                response = HttpResponse(content_type="text/csv; charset=utf-8")
                response["Content-Disposition"] = f'attachment; filename="{nombre_base}_saimyr.csv"'
                response.write("\ufeff")
                writer = csv.writer(response)
                writer.writerow(headers)
                writer.writerows(data)
                return response

            # txt (por defecto)
            response = HttpResponse(content_type="text/plain; charset=utf-8")
            response["Content-Disposition"] = f'attachment; filename="{nombre_base}_saimyr.txt"'
            lines = ["|".join(str(h) for h in headers)]
            for row in data:
                lines.append("|".join("" if v is None else str(v) for v in row))
            response.write("\n".join(lines))
            return response

        # ── Sabaneta: reporte plano de publicidad exterior visual ─────────────
        if proceso.municipio.codigo == "SABANETA":
            from etl.services.exportador_sabaneta import build_rows
            headers, data = build_rows(proceso, {
                "q": q, "q_consec": q_consec, "q_num": q_num, "q_nombre": q_nombre,
                "fecha_desde": fecha_desde, "fecha_hasta": fecha_hasta,
                "estado_pago": estado_pago_filtro,
            })
            if fmt == "csv":
                response = HttpResponse(content_type="text/csv; charset=utf-8")
                response["Content-Disposition"] = f'attachment; filename="{nombre_base}.csv"'
                response.write("﻿")
                writer = csv.writer(response)
                writer.writerow(headers)
                writer.writerows([["" if v is None else v for v in row] for row in data])
                return response
            import openpyxl
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "datos"
            ws.append(headers)
            for row in data:
                ws.append(row)
            buf = io.BytesIO()
            wb.save(buf)
            buf.seek(0)
            resp = HttpResponse(
                buf.read(),
                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            resp["Content-Disposition"] = f'attachment; filename="{nombre_base}.xlsx"'
            return resp

        # ── Envigado: TXT intercalado (encabezado + detalles por consecutivo) ──
        if proceso.municipio.codigo == "ENVIGADO" and fmt == "txt":
            qs_enc = EncabezadoCXC.objects.filter(proceso=proceso).order_by("consecutivo_cxc")
            if q:
                qs_enc = (qs_enc.filter(consecutivo_cxc__icontains=q) |
                          qs_enc.filter(numero_documento__icontains=q) |
                          qs_enc.filter(razon_social__icontains=q))
            if fecha_desde:
                qs_enc = qs_enc.filter(fecha_cobro__gte=fecha_desde)
            if fecha_hasta:
                qs_enc = qs_enc.filter(fecha_cobro__lte=fecha_hasta)
            if estado_filtro:
                qs_enc = qs_enc.filter(estado_cxc__iexact=estado_filtro)
            if estado_pago_filtro:
                qs_enc = qs_enc.filter(estado_pago__icontains=estado_pago_filtro)
            if caldas_periodo_filtro:
                qs_enc = qs_enc.filter(datos_extra__periodo=caldas_periodo_filtro)
            if caldas_ano_filtro:
                qs_enc = qs_enc.filter(datos_extra__ano=caldas_ano_filtro)
            from datetime import datetime as _dt
            for campo, desde, hasta in [
                ("fecha_pago", env_fecha_pago_desde, env_fecha_pago_hasta),
                ("fecha_pres", env_fecha_pres_desde, env_fecha_pres_hasta),
            ]:
                if desde or hasta:
                    vals_en_rango = []
                    for v in EncabezadoCXC.objects.filter(proceso=proceso)\
                            .values_list(f"datos_extra__{campo}", flat=True).distinct():
                        if not v:
                            continue
                        try:
                            d = _dt.strptime(str(v), "%d-%m-%Y").date()
                        except ValueError:
                            continue
                        if desde and d < desde:
                            continue
                        if hasta and d > hasta:
                            continue
                        vals_en_rango.append(v)
                    qs_enc = qs_enc.filter(**{f"datos_extra__{campo}__in": vals_en_rango})

            # Pre-cargar detalles agrupados por consecutivo
            det_map = {}
            for d in DetalleCXC.objects.filter(encabezado__proceso=proceso).select_related("encabezado"):
                det_map.setdefault(d.encabezado.consecutivo_cxc, []).append(d)

            # Mapa codigo_concepto -> descripcion para Envigado
            DESC_ENVIGADO = {
                "8575": "AUTORRETENCION INDUSTRIA Y COMERCIO - INDUSTRIAL",
                "8576": "AUTORRETENCION INDUSTRIA Y COMERCIO - COMERCIAL",
                "8577": "AUTORRETENCION INDUSTRIA Y COMERCIO - SERVICIOS",
                "8578": "SANCIONES",
                "8579": "INTERESES",
                "7377": "Retencion por industria y comercio",
                "7378": "Sancion RETEICA",
                "7379": "interes de mora por RETEICA",
            }

            import unicodedata as _ud
            def _sin_tildes(s):
                return _ud.normalize("NFD", str(s)).encode("ascii", "ignore").decode()

            is_rete = proceso.codigo == "CXC_RETE"
            lines = []
            for enc in qs_enc:
                ex    = enc.datos_extra or {}
                cxc   = enc.consecutivo_cxc or ""
                nit   = enc.numero_documento or ""
                nom   = enc.razon_social or ""
                ano   = ex.get("ano", "")
                per   = ex.get("periodo", "")
                ep    = enc.estado_pago or ""
                fpres = ex.get("fecha_pres", "")
                fpago = ex.get("fecha_pago", "")

                def _val(d):
                    return int(d.valor_total) if d.valor_total is not None else 0

                if is_rete:
                    total_raw = ex.get("total", enc.total_a_pagar or 0)
                    try:
                        total = int(float(total_raw))
                    except (ValueError, TypeError):
                        total = total_raw
                    lines.append(f"{cxc}|{nit}|{per}|{ano}|{nom}|{total}|{ep}|{fpago}|{fpres}")
                    # Agrupar valores por código y descontar exceso del 7377
                    vals_rete = {}
                    exceso = 0
                    for d in det_map.get(cxc, []):
                        cod = d.codigo_concepto or ""
                        v   = _val(d)
                        if not cod:  # RETENCION PRACTICADA EN EXCESO (negativo)
                            exceso += abs(v)
                        else:
                            vals_rete[cod] = vals_rete.get(cod, 0) + v
                    # Restar exceso del 7377
                    if exceso and "7377" in vals_rete:
                        vals_rete["7377"] = vals_rete["7377"] - exceso
                    for cod, v in vals_rete.items():
                        if v <= 0:
                            continue
                        desc = _sin_tildes(DESC_ENVIGADO.get(cod, cod))
                        if not desc:
                            continue
                        lines.append(f"{cxc}|{nit}|{nom}|{cod}|{desc}|{v}")
                else:
                    try:
                        total_auto = int(float(enc.total_a_pagar)) if enc.total_a_pagar is not None else ""
                    except (ValueError, TypeError):
                        total_auto = enc.total_a_pagar or ""
                    lines.append(f"{cxc}|{nit}|{per}|{ano}|{nom}|{total_auto}|{ep}|{fpago}|{fpres}")
                    for d in det_map.get(cxc, []):
                        v = _val(d)
                        if not v:
                            continue
                        cod  = d.codigo_concepto or ""
                        desc = _sin_tildes(DESC_ENVIGADO.get(cod, cod))
                        lines.append(f"{cxc}|{nit}|{nom}|{cod}|{desc}|{v}")

            suffix = "rete" if is_rete else "auto"
            response = HttpResponse(content_type="text/plain; charset=utf-8")
            response["Content-Disposition"] = f'attachment; filename="{nombre_base}_{suffix}.txt"'
            response.write("\n".join(lines))
            return response

        # ── Envigado: Excel / CSV con columnas del dashboard ─────────────────
        if proceso.municipio.codigo == "ENVIGADO" and fmt in ("excel", "csv"):
            DESC_ENVIGADO = {
                "8575": "AUTORRETENCIÓN INDUSTRIA Y COMERCIO - INDUSTRIAL",
                "8576": "AUTORRETENCIÓN INDUSTRIA Y COMERCIO - COMERCIAL",
                "8577": "AUTORRETENCION INDUSTRIA Y COMERCIO - SERVICIOS",
                "8578": "SANCIONES",
                "8579": "INTERESES",
                "7377": "Retención por industria y comercio",
                "7378": "Sanción RETEICA",
                "7379": "interes de mora por RETEICA",
            }
            # Aplicar filtros al encabezado
            qs_enc = EncabezadoCXC.objects.filter(proceso=proceso).order_by("consecutivo_cxc")
            if q:
                qs_enc = (qs_enc.filter(consecutivo_cxc__icontains=q) |
                          qs_enc.filter(numero_documento__icontains=q) |
                          qs_enc.filter(razon_social__icontains=q))
            if q_consec:
                qs_enc = qs_enc.filter(consecutivo_cxc__icontains=q_consec)
            if q_num:
                qs_enc = qs_enc.filter(numero_documento__icontains=q_num)
            if q_nombre:
                qs_enc = qs_enc.filter(razon_social__icontains=q_nombre)
            if estado_pago_filtro:
                qs_enc = qs_enc.filter(estado_pago__icontains=estado_pago_filtro)
            if caldas_periodo_filtro:
                qs_enc = qs_enc.filter(datos_extra__periodo=caldas_periodo_filtro)
            if caldas_ano_filtro:
                qs_enc = qs_enc.filter(datos_extra__ano=caldas_ano_filtro)
            from datetime import datetime as _dt5
            for _campo, _desde, _hasta in [
                ("fecha_pago", env_fecha_pago_desde, env_fecha_pago_hasta),
                ("fecha_pres", env_fecha_pres_desde, env_fecha_pres_hasta),
            ]:
                if _desde or _hasta:
                    _vals = []
                    for _v in EncabezadoCXC.objects.filter(proceso=proceso)\
                            .values_list(f"datos_extra__{_campo}", flat=True).distinct():
                        if not _v:
                            continue
                        try:
                            _d = _dt5.strptime(str(_v), "%d-%m-%Y").date()
                        except ValueError:
                            continue
                        if _desde and _d < _desde:
                            continue
                        if _hasta and _d > _hasta:
                            continue
                        _vals.append(_v)
                    qs_enc = qs_enc.filter(**{f"datos_extra__{_campo}__in": _vals})

            enc_hdrs = ["Consecutivo", "NIT", "Nombre / Razón Social",
                        "Periodo", "Año", "Fecha Presentación", "Fecha Pago",
                        "Total a Pagar", "Estado Pago"]
            from datetime import datetime as _dt6

            def _parse_dmY(v):
                """'DD-MM-YYYY' → datetime.date o None"""
                try:
                    return _dt6.strptime(str(v), "%d-%m-%Y").date()
                except Exception:
                    return None

            def _fmt_fecha(v):
                """Para CSV: 'DD/MM/YYYY'"""
                d = _parse_dmY(v)
                return d.strftime("%d/%m/%Y") if d else (str(v) if v else "")

            enc_rows = []
            for enc in qs_enc:
                ex = enc.datos_extra or {}
                enc_rows.append([
                    enc.consecutivo_cxc,
                    enc.numero_documento,
                    enc.razon_social,
                    ex.get("periodo", ""),
                    ex.get("ano", ""),
                    ex.get("fecha_pres", ""),   # se reemplaza por date en Excel
                    ex.get("fecha_pago", ""),   # se reemplaza por date en Excel
                    enc.total_a_pagar,
                    enc.estado_pago,
                ])

            # Detalles filtrados por los mismos encabezados
            enc_ids = list(qs_enc.values_list("id", flat=True))
            qs_det = DetalleCXC.objects.filter(encabezado_id__in=enc_ids)\
                         .select_related("encabezado").order_by("encabezado__consecutivo_cxc")
            if q_concepto:
                qs_det = qs_det.filter(codigo_concepto__icontains=q_concepto)
            det_hdrs = ["Consecutivo CXC", "NIT", "Nombre", "Código Concepto", "Descripción", "Valor Total"]
            det_rows = []
            for d in qs_det:
                det_rows.append([
                    d.encabezado.consecutivo_cxc,
                    d.encabezado.numero_documento,
                    d.encabezado.razon_social,
                    d.codigo_concepto,
                    DESC_ENVIGADO.get(d.codigo_concepto, d.codigo_concepto),
                    int(d.valor_total) if d.valor_total is not None else 0,
                ])

            # Índices de columnas de fecha en enc_rows (0-based): 5=fecha_pres, 6=fecha_pago
            DATE_COLS_ENC = {5, 6}

            if fmt == "excel":
                import openpyxl
                from openpyxl.styles import Font, PatternFill, Alignment
                wb = openpyxl.Workbook()
                for sheet_name, hdrs, rows, date_cols in [
                    ("Encabezado", enc_hdrs, enc_rows, DATE_COLS_ENC),
                    ("Detalle",    det_hdrs, det_rows, set()),
                ]:
                    ws = wb.create_sheet(sheet_name)
                    ws.append(hdrs)
                    for cell in ws[1]:
                        cell.font = Font(bold=True, color="FFFFFF")
                        cell.fill = PatternFill("solid", fgColor="1a6b3c")
                        cell.alignment = Alignment(horizontal="center")
                    for row in rows:
                        # Convertir strings de fecha a date objects para las columnas de fecha
                        converted = []
                        for i, val in enumerate(row):
                            if i in date_cols:
                                d = _parse_dmY(val)
                                converted.append(d if d else val)
                            else:
                                converted.append(val)
                        ws.append(converted)
                    # Aplicar formato de fecha DD/MM/YYYY a columnas de fecha
                    for col_idx in date_cols:
                        col_letter = openpyxl.utils.get_column_letter(col_idx + 1)
                        for cell in ws[col_letter][1:]:  # skip header
                            cell.number_format = "DD/MM/YYYY"
                del wb["Sheet"]
                buf = io.BytesIO()
                wb.save(buf)
                buf.seek(0)
                suffix = "rete" if proceso.codigo == "CXC_RETE" else "auto"
                resp = HttpResponse(
                    buf.read(),
                    content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )
                resp["Content-Disposition"] = f'attachment; filename="{nombre_base}_{suffix}.xlsx"'
                return resp

            # CSV: solo el tab activo — fechas como DD/MM/YYYY
            response = HttpResponse(content_type="text/csv; charset=utf-8")
            suffix = "rete" if proceso.codigo == "CXC_RETE" else "auto"
            response["Content-Disposition"] = f'attachment; filename="{nombre_base}_{suffix}.csv"'
            response.write("\ufeff")
            writer = csv.writer(response)
            if tab == "detalle":
                writer.writerow(det_hdrs)
                writer.writerows(det_rows)
            else:
                writer.writerow(enc_hdrs)
                csv_rows = []
                for row in enc_rows:
                    csv_row = list(row)
                    for i in DATE_COLS_ENC:
                        csv_row[i] = _fmt_fecha(csv_row[i])
                    csv_rows.append(csv_row)
                writer.writerows(csv_rows)
            return response

        # ── helpers para construir los dos datasets ──────────────────────────
        def _enc_qs():
            qs = EncabezadoCXC.objects.filter(proceso=proceso)
            if q:
                qs = (qs.filter(consecutivo_cxc__icontains=q) | qs.filter(numero_documento__icontains=q) |
                      qs.filter(primer_apellido__icontains=q) | qs.filter(razon_social__icontains=q) |
                      qs.filter(descripcion__icontains=q))
            if q_consec:
                qs = qs.filter(consecutivo_cxc__icontains=q_consec)
            if q_num:
                qs = qs.filter(numero_documento__icontains=q_num)
            if q_nombre:
                qs = (qs.filter(primer_apellido__icontains=q_nombre) |
                      qs.filter(razon_social__icontains=q_nombre))
            if q_desc:
                qs = qs.filter(descripcion__icontains=q_desc)
            if tipo_doc_filtro:
                qs = qs.filter(tipo_documento__iexact=tipo_doc_filtro)
            if estado_pago_filtro:
                qs = qs.filter(estado_pago__icontains=estado_pago_filtro)
            if fecha_desde:
                qs = qs.filter(fecha_cobro__gte=fecha_desde)
            if fecha_hasta:
                qs = qs.filter(fecha_cobro__lte=fecha_hasta)
            if estado_filtro == "__blank__":
                qs = qs.filter(estado_cxc="")
            elif estado_filtro:
                qs = qs.filter(estado_cxc__iexact=estado_filtro)
            if caldas_tipo_persona_filtro:
                qs = qs.filter(datos_extra__tipo_persona=caldas_tipo_persona_filtro)
            if caldas_periodo_filtro:
                qs = qs.filter(datos_extra__periodo=caldas_periodo_filtro)
            if caldas_ano_filtro:
                qs = qs.filter(datos_extra__ano=caldas_ano_filtro)
            if q_total_pagar:
                from django.db.models.functions import Cast as _Cast
                from django.db.models import CharField as _CF
                qs = qs.annotate(_tp_s=_Cast("total_a_pagar", _CF())).filter(_tp_s__icontains=q_total_pagar)
            from datetime import datetime as _dt4
            for _campo, _desde, _hasta in [
                ("fecha_pago", env_fecha_pago_desde, env_fecha_pago_hasta),
                ("fecha_pres", env_fecha_pres_desde, env_fecha_pres_hasta),
            ]:
                if _desde or _hasta:
                    _vals = []
                    for _v in EncabezadoCXC.objects.filter(proceso=proceso)\
                            .values_list(f"datos_extra__{_campo}", flat=True).distinct():
                        if not _v:
                            continue
                        try:
                            _d = _dt4.strptime(str(_v), "%d-%m-%Y").date()
                        except ValueError:
                            continue
                        if _desde and _d < _desde:
                            continue
                        if _hasta and _d > _hasta:
                            continue
                        _vals.append(_v)
                    qs = qs.filter(**{f"datos_extra__{_campo}__in": _vals})
            hdrs = ["consecutivo_cxc", "tipo_documento", "numero_documento",
                    "primer_nombre", "segundo_nombre", "primer_apellido", "segundo_apellido",
                    "razon_social", "fecha_cobro", "fecha_vencimiento", "descripcion",
                    "total_a_pagar", "estado_pago", "estado_cxc"]
            return hdrs, list(qs.values_list(*hdrs))

        def _det_qs():
            qs = DetalleCXC.objects.filter(encabezado__proceso=proceso).select_related("encabezado")
            # Aplicar todos los filtros de encabezado (filtro inteligente cruzado)
            if q:
                qs = (qs.filter(codigo_concepto__icontains=q) |
                      qs.filter(encabezado__consecutivo_cxc__icontains=q) |
                      qs.filter(centro_costo__icontains=q))
            if q_consec:
                qs = qs.filter(encabezado__consecutivo_cxc__icontains=q_consec)
            if q_num:
                qs = qs.filter(encabezado__numero_documento__icontains=q_num)
            if q_nombre:
                qs = (qs.filter(encabezado__primer_apellido__icontains=q_nombre) |
                      qs.filter(encabezado__razon_social__icontains=q_nombre))
            if q_desc:
                qs = qs.filter(encabezado__descripcion__icontains=q_desc)
            if tipo_doc_filtro:
                qs = qs.filter(encabezado__tipo_documento__iexact=tipo_doc_filtro)
            if estado_pago_filtro:
                qs = qs.filter(encabezado__estado_pago__icontains=estado_pago_filtro)
            if fecha_desde:
                qs = qs.filter(encabezado__fecha_cobro__gte=fecha_desde)
            if fecha_hasta:
                qs = qs.filter(encabezado__fecha_cobro__lte=fecha_hasta)
            if estado_filtro == "__blank__":
                qs = qs.filter(encabezado__estado_cxc="")
            elif estado_filtro:
                qs = qs.filter(encabezado__estado_cxc__iexact=estado_filtro)
            if caldas_tipo_persona_filtro:
                qs = qs.filter(encabezado__datos_extra__tipo_persona=caldas_tipo_persona_filtro)
            if caldas_periodo_filtro:
                qs = qs.filter(encabezado__datos_extra__periodo=caldas_periodo_filtro)
            if caldas_ano_filtro:
                qs = qs.filter(encabezado__datos_extra__ano=caldas_ano_filtro)
            if q_concepto:
                qs = qs.filter(codigo_concepto__icontains=q_concepto)
            if q_centro or caldas_clasi_filtro:
                cc = caldas_clasi_filtro or q_centro
                qs = qs.filter(centro_costo__icontains=cc)
            if q_valor_unit:
                from django.db.models.functions import Cast as _Cast
                from django.db.models import CharField as _CF
                qs = qs.annotate(_vu_s=_Cast("valor_unitario", _CF())).filter(_vu_s__icontains=q_valor_unit)
            if q_valor_tot:
                from django.db.models.functions import Cast as _Cast
                from django.db.models import CharField as _CF
                qs = qs.annotate(_vt_s=_Cast("valor_total", _CF())).filter(_vt_s__icontains=q_valor_tot)
            from datetime import datetime as _dt3
            for _campo, _desde, _hasta in [
                ("fecha_pago", env_fecha_pago_desde, env_fecha_pago_hasta),
                ("fecha_pres", env_fecha_pres_desde, env_fecha_pres_hasta),
            ]:
                if _desde or _hasta:
                    _vals = []
                    for _v in EncabezadoCXC.objects.filter(proceso=proceso)\
                            .values_list(f"datos_extra__{_campo}", flat=True).distinct():
                        if not _v:
                            continue
                        try:
                            _d = _dt3.strptime(str(_v), "%d-%m-%Y").date()
                        except ValueError:
                            continue
                        if _desde and _d < _desde:
                            continue
                        if _hasta and _d > _hasta:
                            continue
                        _vals.append(_v)
                    qs = qs.filter(**{f"encabezado__datos_extra__{_campo}__in": _vals})
            rows = list(qs.values("encabezado__consecutivo_cxc", "codigo_concepto",
                                  "centro_costo", "cantidad", "valor_unitario", "valor_total"))
            hdrs = ["consecutivo_cxc", "codigo_concepto", "centro_costo", "cantidad", "valor_unitario", "valor_total"]
            data = [[r["encabezado__consecutivo_cxc"], r["codigo_concepto"], r["centro_costo"],
                     r["cantidad"], r["valor_unitario"], r["valor_total"]] for r in rows]
            return hdrs, data

        # ── Excel: dos hojas ─────────────────────────────────────────────────
        if fmt == "excel":
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment
            wb = openpyxl.Workbook()

            enc_hdrs, enc_data = _enc_qs()
            det_hdrs, det_data = _det_qs()

            for sheet_name, hdrs, data in [("Encabezado", enc_hdrs, enc_data),
                                            ("Detalle", det_hdrs, det_data)]:
                ws = wb.active if sheet_name == "Encabezado" else wb.create_sheet(sheet_name)
                ws.title = sheet_name
                ws.append(hdrs)
                # estilo encabezado
                for cell in ws[1]:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="1a6b3c")
                    cell.alignment = Alignment(horizontal="center")
                for row in data:
                    ws.append(["" if v is None else v for v in row])

            buf = io.BytesIO()
            wb.save(buf)
            buf.seek(0)
            response = HttpResponse(
                buf.read(),
                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
            response["Content-Disposition"] = f'attachment; filename="{nombre_base}.xlsx"'
            return response

        # ── CSV / TXT: pestaña activa ────────────────────────────────────────
        if tab == "detalle":
            headers, data = _det_qs()
        else:
            headers, data = _enc_qs()

        nombre = f"{nombre_base}_{tab}"

        if fmt == "csv":
            response = HttpResponse(content_type="text/csv; charset=utf-8")
            response["Content-Disposition"] = f'attachment; filename="{nombre}.csv"'
            response.write("\ufeff")
            writer = csv.writer(response)
            writer.writerow(headers)
            writer.writerows(data)
            return response

        else:  # txt
            response = HttpResponse(content_type="text/plain; charset=utf-8")
            response["Content-Disposition"] = f'attachment; filename="{nombre}.txt"'
            lines = ["|".join(str(h) for h in headers)]
            for row in data:
                lines.append("|".join("" if v is None else str(v) for v in row))
            response.write("\n".join(lines))
            return response
